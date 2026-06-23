"""U-Net 训练入口。

该文件负责完成一条完整的语义分割训练流水线：

1. 从 ``data/imgs`` 和 ``data/masks`` 读取图像与标注；
2. 将数据划分为训练集和验证集；
3. 创建 U-Net、优化器、损失函数和学习率调度器；
4. 逐批执行前向传播、反向传播和参数更新；
5. 每个 epoch 结束后计算 Dice、IoU、Precision、Recall 和 HD95；
6. 将每轮训练得到的模型权重保存到 ``checkpoints``。

运行示例：
    python train.py --epochs 5 --batch-size 2 --amp
"""

import argparse
import logging
import os
import random
import sys

import torch
import torch.nn as nn
import torch.nn.functional as F
import torchvision.transforms as transforms
import torchvision.transforms.functional as TF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pathlib import Path
from torch import optim
from torch.utils.data import DataLoader, random_split
from tqdm import tqdm
import wandb
from evaluate import evaluate
from unet import UNet
from utils.data_loading import BasicDataset, CarvanaDataset
from utils.dice_score import dice_loss

# 项目约定的数据与模型保存位置。
# Path 对象可以跨平台处理路径，后续也便于创建不存在的目录。
dir_img = Path('./data/imgs/')
dir_mask = Path('./data/masks/')
dir_checkpoint = Path('./checkpoints/')
metrics_file = Path('./training_metrics.xlsx')


def create_metrics_workbook(path: Path):
    """创建新的训练指标工作簿，覆盖同名旧文件。"""

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Training Metrics'

    headers = [
        'Epoch',
        'Train Loss',
        'Dice',
        'IoU',
        'Precision',
        'Recall',
        'HD95',
        'Learning Rate',
    ]
    worksheet.append(headers)

    header_fill = PatternFill(fill_type='solid', fgColor='1F4E78')
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    column_widths = {
        'A': 10,
        'B': 16,
        'C': 12,
        'D': 12,
        'E': 14,
        'F': 12,
        'G': 12,
        'H': 16,
    }
    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width

    worksheet.freeze_panes = 'A2'
    worksheet.auto_filter.ref = 'A1:H1'
    workbook.save(path)
    return workbook, worksheet


def append_epoch_metrics(workbook, worksheet, path: Path, epoch: int, loss: float, metrics: dict, lr: float):
    """追加一个 epoch 的汇总指标并立即保存。"""

    worksheet.append([
        epoch,
        loss,
        metrics['dice'],
        metrics['iou'],
        metrics['precision'],
        metrics['recall'],
        metrics['hd95'],
        lr,
    ])

    row = worksheet.max_row
    worksheet.cell(row, 1).number_format = '0'
    for column in range(2, 8):
        worksheet.cell(row, column).number_format = '0.000000'
    worksheet.cell(row, 8).number_format = '0.00000000E+00'
    worksheet.auto_filter.ref = f'A1:H{row}'
    workbook.save(path)


def train_model(
        model,
        device,
        epochs: int = 5,
        batch_size: int = 1,
        learning_rate: float = 1e-5,
        val_percent: float = 0.1,
        save_checkpoint: bool = True,
        img_scale: float = 0.5,
        amp: bool = False,
        weight_decay: float = 1e-8,
        momentum: float = 0.999,
        gradient_clipping: float = 1.0,
):
    """训练一个已经创建好的 U-Net 模型。

    参数:
        model: 待训练的 U-Net 实例。
        device: 计算设备，例如 ``cuda``、``cpu`` 或 ``mps``。
        epochs: 完整遍历训练集的次数。
        batch_size: 每次参数更新使用的样本数量。
        learning_rate: 优化器的初始学习率。
        val_percent: 验证集占全部数据的比例，取值范围为 0～1。
        save_checkpoint: 是否在每个 epoch 结束后保存模型权重。
        img_scale: 输入图像的缩放比例，可用于降低显存占用。
        amp: 是否启用自动混合精度训练。
        weight_decay: RMSprop 的权重衰减系数，用于抑制过拟合。
        momentum: RMSprop 的动量参数。
        gradient_clipping: 梯度裁剪上限，用于降低梯度爆炸风险。
    """

    # 1. 创建数据集。
    # Carvana 数据集的掩码文件名带有 ``_mask`` 后缀，因此优先按该规则加载；
    # 如果目录结构不符合 Carvana 规则，则退回到图像和掩码同名的通用数据集。
    try:
        dataset = CarvanaDataset(dir_img, dir_mask, img_scale)
    except (AssertionError, RuntimeError, IndexError):
        dataset = BasicDataset(dir_img, dir_mask, img_scale)

    # 2. 划分训练集和验证集。
    # 固定随机种子可以保证每次运行时得到相同的数据划分，方便复现实验结果。
    n_val = int(len(dataset) * val_percent)
    n_train = len(dataset) - n_val
    train_set, val_set = random_split(dataset, [n_train, n_val], generator=torch.Generator().manual_seed(0))

    # 3. 创建数据加载器。
    # num_workers 使用 CPU 核心数并行读取数据；pin_memory 可加快 CPU 到 CUDA 的数据拷贝。
    loader_args = dict(batch_size=batch_size, num_workers=os.cpu_count(), pin_memory=True)
    train_loader = DataLoader(train_set, shuffle=True, **loader_args)
    # 验证集不需要打乱，也不能丢弃最后一个不足 batch_size 的批次。
    val_loader = DataLoader(val_set, shuffle=False, drop_last=False, **loader_args)

    # 初始化 Weights & Biases 实验，用来记录损失、验证指标和模型参数分布。
    # anonymous='must' 允许在没有登录账号时以匿名方式运行。
    experiment = wandb.init(project='U-Net', resume='allow', anonymous='must')
    experiment.config.update(
        dict(epochs=epochs, batch_size=batch_size, learning_rate=learning_rate,
             val_percent=val_percent, save_checkpoint=save_checkpoint, img_scale=img_scale, amp=amp)
    )

    # 每次启动新训练时覆盖旧指标文件，防止不同实验的数据混在一起。
    metrics_workbook, metrics_worksheet = create_metrics_workbook(metrics_file)

    logging.info(f'''Starting training:
        Epochs:          {epochs}
        Batch size:      {batch_size}
        Learning rate:   {learning_rate}
        Training size:   {n_train}
        Validation size: {n_val}
        Checkpoints:     {save_checkpoint}
        Device:          {device.type}
        Images scaling:  {img_scale}
        Mixed Precision: {amp}
    ''')

    # 4. 配置优化器、学习率调度器、混合精度缩放器和损失函数。
    # RMSprop 是原始 U-Net 相关实现中常用的优化器；foreach=True 可批量处理参数以提高效率。
    optimizer = optim.RMSprop(model.parameters(),
                              lr=learning_rate, weight_decay=weight_decay, momentum=momentum, foreach=True)
    # 当验证 Dice 分数连续若干次没有提高时降低学习率；mode='max' 表示分数越高越好。
    scheduler = optim.lr_scheduler.ReduceLROnPlateau(
                            optimizer,
                            mode='max',
                            factor=0.5,
                            patience=7,
                            min_lr=1e-7
)    
    # GradScaler 在 AMP 模式下放大损失，减少 float16 梯度下溢；未启用 AMP 时相当于普通训练。
    grad_scaler = torch.cuda.amp.GradScaler(enabled=amp)
    # 多分类分割使用交叉熵；单类别前景/背景分割使用二元交叉熵。
    criterion = nn.CrossEntropyLoss() if model.n_classes > 1 else nn.BCEWithLogitsLoss()
    # global_step 记录已处理的训练批次数，用作 WandB 日志横轴。
    global_step = 0

    # 5. 开始训练。
    for epoch in range(1, epochs + 1):
        model.train()
        epoch_loss = 0
        train_progress = tqdm(
            train_loader,
            desc=f'Epoch {epoch}/{epochs}',
            unit='batch',
        )
        for batch in train_progress:
            # DataLoader 返回字典；图像形状通常为 [N, C, H, W]，
            # 掩码形状通常为 [N, H, W]。
            images, true_masks = batch['image'], batch['mask']

            # 在训练开始处检查通道数，可尽早发现灰度/RGB 图像配置不匹配的问题。
            assert images.shape[1] == model.n_channels, \
                f'Network has been defined with {model.n_channels} input channels, ' \
                f'but loaded images have {images.shape[1]} channels. Please check that ' \
                'the images are loaded correctly.'

            # channels_last 在部分 GPU 卷积场景下具有更好的内存访问效率。
            images = images.to(device=device, dtype=torch.float32, memory_format=torch.channels_last)
            # CrossEntropyLoss 要求类别标签使用整数索引，因此掩码转为 long。
            true_masks = true_masks.to(device=device, dtype=torch.long)

            # autocast 会根据设备自动选择较低精度执行适合的算子，从而减少显存并提升速度。
            with torch.autocast(device.type if device.type != 'mps' else 'cpu', enabled=amp):
                # masks_pred 为未经归一化的 logits。
                # 多分类形状为 [N, classes, H, W]；二分类形状为 [N, 1, H, W]。
                masks_pred = model(images)
                if model.n_classes == 1:
                    # 单输出通道时去掉通道维，并组合 BCE 与 Dice 损失。
                    bce_loss = criterion(masks_pred.squeeze(1),true_masks.float())
                    d_loss = dice_loss(torch.sigmoid(masks_pred.squeeze(1)),true_masks.float(),multiclass=False)
                    loss = 0.4 * bce_loss + 0.6 * d_loss
                else:
                    # 多分类时，交叉熵直接接收 logits 和类别索引。
                    loss = criterion(masks_pred, true_masks)
                    # Dice 损失需要概率图和 one-hot 标签，维度调整为 [N, C, H, W]。
                    loss += dice_loss(
                        F.softmax(masks_pred, dim=1).float(),
                        F.one_hot(true_masks, model.n_classes).permute(0, 3, 1, 2).float(),
                        multiclass=True
                    )

            # set_to_none=True 比把梯度清零更节省内存，并允许 PyTorch 跳过部分无梯度参数。
            optimizer.zero_grad(set_to_none=True)
            # AMP 下依次执行：缩放损失、反向传播、还原梯度、裁剪梯度、更新参数。
            grad_scaler.scale(loss).backward()
            grad_scaler.unscale_(optimizer)
            torch.nn.utils.clip_grad_norm_(model.parameters(), gradient_clipping)
            grad_scaler.step(optimizer)
            grad_scaler.update()

            global_step += 1
            # loss 默认是当前批次的均值；乘以批次图片数后累加，便于得到严格的 epoch 样本均值。
            epoch_loss += loss.item() * images.shape[0]
            train_progress.set_postfix(loss=f'{loss.item():.4f}')
            experiment.log({
                'train loss': loss.item(),
                'step': global_step,
                'epoch': epoch
            })

        average_epoch_loss = epoch_loss / max(n_train, 1)
        # 每个 epoch 只进行一次完整验证。
        val_metrics = evaluate(model, val_loader, device, amp, show_progress=False)
        scheduler.step(val_metrics['dice'])

        # 每轮只收集一次参数和梯度直方图。
        histograms = {}
        for tag, value in model.named_parameters():
            tag = tag.replace('/', '.')
            if not (torch.isinf(value) | torch.isnan(value)).any():
                histograms['Weights/' + tag] = wandb.Histogram(value.data.cpu())
            if value.grad is not None and not (torch.isinf(value.grad) | torch.isnan(value.grad)).any():
                histograms['Gradients/' + tag] = wandb.Histogram(value.grad.data.cpu())

        epoch_log = {
            'epoch average train loss': average_epoch_loss,
            'learning rate': optimizer.param_groups[0]['lr'],
            'validation Dice': val_metrics['dice'],
            'validation IoU': val_metrics['iou'],
            'validation Precision': val_metrics['precision'],
            'validation Recall': val_metrics['recall'],
            'validation HD95': val_metrics['hd95'],
            'epoch': epoch,
            'step': global_step,
            **histograms,
        }
        try:
            pred_example = (
                (torch.sigmoid(masks_pred).squeeze(1) > 0.5).float()
                if model.n_classes == 1
                else masks_pred.argmax(dim=1).float()
            )
            epoch_log.update({
                'images': wandb.Image(images[0].cpu()),
                'masks': {
                    'true': wandb.Image(true_masks[0].float().cpu()),
                    'pred': wandb.Image(pred_example[0].cpu()),
                },
            })
        except Exception:
            # 可视化日志失败不应中断模型训练。
            pass
        experiment.log(epoch_log)

        # 每个 epoch 结束后保存 state_dict，而不是整个模型对象，
        # 这样权重文件体积更小，也不依赖保存时的 Python 对象结构。
        checkpoint_status = 'not saved'
        if save_checkpoint:
            Path(dir_checkpoint).mkdir(parents=True, exist_ok=True)
            state_dict = model.state_dict()
            # 同时保存掩码中的原始像素值，预测阶段可将类别索引还原为原始颜色/灰度值。
            state_dict['mask_values'] = dataset.mask_values
            torch.save(state_dict, str(dir_checkpoint / 'checkpoint_epoch{}.pth'.format(epoch)))
            checkpoint_status = 'saved'

        append_epoch_metrics(
            metrics_workbook,
            metrics_worksheet,
            metrics_file,
            epoch,
            average_epoch_loss,
            val_metrics,
            optimizer.param_groups[0]['lr'],
        )

        # 每个 epoch 的控制台输出严格合并为一条日志。
        logging.info(
            'Epoch %d/%d | loss: %.6f | Dice: %.4f | IoU: %.4f | Precision: %.4f | '
            'Recall: %.4f | HD95: %.4f | LR: %.8g | checkpoint: %s',
            epoch,
            epochs,
            average_epoch_loss,
            val_metrics['dice'],
            val_metrics['iou'],
            val_metrics['precision'],
            val_metrics['recall'],
            val_metrics['hd95'],
            optimizer.param_groups[0]['lr'],
            checkpoint_status,
        )


def get_args():
    """定义并解析命令行参数。"""

    parser = argparse.ArgumentParser(description='Train the UNet on images and target masks')
    parser.add_argument('--epochs', '-e', metavar='E', type=int, default=100, help='Number of epochs')
    parser.add_argument('--batch-size', '-b', dest='batch_size', metavar='B', type=int, default=8, help='Batch size')
    parser.add_argument('--learning-rate', '-l', metavar='LR', type=float, default=1e-4,
                        help='Learning rate', dest='lr')
    parser.add_argument('--load', '-f', type=str, default=False, help='Load model from a .pth file')
    parser.add_argument('--scale', '-s', type=float, default=1, help='Downscaling factor of the images')
    parser.add_argument('--validation', '-v', dest='val', type=float, default=10.0,
                        help='Percent of the data that is used as validation (0-100)')
    parser.add_argument('--amp', action='store_true', default=False, help='Use mixed precision')
    parser.add_argument('--bilinear', action='store_true', default=False, help='Use bilinear upsampling')
    parser.add_argument('--classes', '-c', type=int, default=1, help='Number of classes')

    return parser.parse_args()


if __name__ == '__main__':
    # 只有直接执行本文件时才开始训练；作为模块导入时不会触发。
    args = get_args()

    logging.basicConfig(level=logging.INFO, format='%(levelname)s: %(message)s')
    # 优先使用 NVIDIA CUDA；不可用时回退到 CPU。
    device = torch.device('cuda' if torch.cuda.is_available() else 'cpu')
    logging.info(f'Using device {device}')

    # 根据任务创建网络：
    # n_channels=3 表示输入为 RGB 图像；
    # n_classes 表示每个像素需要预测的类别数量。
    model = UNet(n_channels=3, n_classes=args.classes, bilinear=args.bilinear)
    # 将卷积网络内部张量优先设置为 channels_last 内存布局。
    model = model.to(memory_format=torch.channels_last)

    logging.info(f'Network:\n'
                 f'\t{model.n_channels} input channels\n'
                 f'\t{model.n_classes} output channels (classes)\n'
                 f'\t{"Bilinear" if model.bilinear else "Transposed conv"} upscaling')

    if args.load:
        # map_location 确保即使权重在 GPU 上保存，也可在当前选择的设备上加载。
        state_dict = torch.load(args.load, map_location=device)
        # mask_values 是数据元信息，不属于模型参数，加载权重前需要移除。
        del state_dict['mask_values']
        model.load_state_dict(state_dict)
        logging.info(f'Model loaded from {args.load}')

    model.to(device=device)
    try:
        # 首次按用户参数启动训练。
        train_model(
            model=model,
            epochs=args.epochs,
            batch_size=args.batch_size,
            learning_rate=args.lr,
            device=device,
            img_scale=args.scale,
            val_percent=args.val / 100,
            amp=args.amp
        )
    except torch.cuda.OutOfMemoryError:
        # 如果显存不足，则清理缓存并启用梯度检查点。
        # 梯度检查点通过在反向传播时重新计算部分前向结果来降低显存占用。
        logging.error('Detected OutOfMemoryError! '
                      'Enabling checkpointing to reduce memory usage, but this slows down training. '
                      'Consider enabling AMP (--amp) for fast and memory efficient training')
        torch.cuda.empty_cache()
        model.use_checkpointing()
        train_model(
            model=model,
            epochs=args.epochs,
            batch_size=args.batch_size,
            learning_rate=args.lr,
            device=device,
            img_scale=args.scale,
            val_percent=args.val / 100,
            amp=args.amp
        )
